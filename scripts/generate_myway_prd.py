#!/usr/bin/env python3
from __future__ import annotations

from datetime import date
from pathlib import Path
import json

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC_PATH = ROOT / "docs" / "MyWay-Product-Requirements-Document.docx"
DIAGRAM_PATH = ROOT / "docs" / "myway-aws-architecture.png"
QUESTIONNAIRE_JSON = ROOT / "docs" / "questionnaire-quality-results.json"


def add_page_number(run):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_toc(paragraph):
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \\o "1-3" \\h \\z \\u'

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "Right-click and choose Update Field in Microsoft Word to build the table of contents."

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r = paragraph.add_run()._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text)
    r.append(fld_end)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.runs[0]
    r.italic = True
    r.font.size = Pt(10)


def draw_architecture(path: Path):
    fig, ax = plt.subplots(figsize=(16, 10), dpi=220)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")

    COLORS = {
        "app": "#DDEBFF",
        "ai": "#EFE3FF",
        "storage": "#E7F7E7",
        "optional": "#FFF3D6",
    }

    def box(x, y, w, h, label, color, fs=12):
        patch = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=2", linewidth=1.2, edgecolor="#333333", facecolor=color)
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=fs, wrap=True)
        return (x + w / 2, y + h / 2)

    def arrow(p1, p2, style="->", dashed=False, lw=1.8):
        ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle=style, mutation_scale=12, linewidth=lw, linestyle="--" if dashed else "-", color="#444444"))

    browser = box(3, 72, 14, 12, "Student Browser\n(Next.js UI)", COLORS["app"])
    nextjs = box(22, 66, 24, 24, "MyWay Next.js App\nRoutes + API Handlers\n/api/resume, /api/analysis\n/api/analysis/intake, /api/analysis/swap\n/api/auth/session, /api/jobs", COLORS["app"], fs=11)
    cognito = box(3, 52, 14, 12, "Amazon Cognito\nSRP + Email Confirm\nOptional TOTP MFA", COLORS["storage"], fs=11)
    bedrock = box(51, 68, 18, 16, "Amazon Bedrock\nClaude Haiku 4.5\nConverse + tool use\nTitan Embeddings V2", COLORS["ai"], fs=11)
    resume_store = box(72, 70, 24, 12, "S3 (resume bytes)\n+DynamoDB resumes\n(metadata)", COLORS["storage"], fs=11)
    analyses = box(72, 54, 24, 12, "DynamoDB analyses\nprofile/intake/embedding\nplan/advisor/improver\nrouting/swapper", COLORS["storage"], fs=10)

    evt = box(25, 40, 18, 10, "EventBridge\nSchedules", COLORS["optional"], fs=11)
    jobs_lambda = box(46, 36, 16, 14, "Lambda\njobs-scraper\n(12-hour)", COLORS["optional"], fs=11)
    courses_lambda = box(46, 18, 16, 14, "Lambda\ncourses-scraper\n(weekly)", COLORS["optional"], fs=11)
    market_s3 = box(67, 26, 15, 16, "S3 snapshots\njobs/latest.json\ncourses/latest.json", COLORS["storage"], fs=10)
    mcf = box(85, 36, 12, 14, "MyCareersFuture\nPublic jobs API", COLORS["optional"], fs=10)
    ssg = box(85, 18, 12, 14, "SSG-WSG APIs\nSkills Framework\nCourse directory", COLORS["optional"], fs=10)

    tf = box(3, 16, 14, 12, "Terraform\n(iac/*.tf)", COLORS["storage"], fs=11)
    gha = box(22, 16, 18, 12, "GitHub Actions\nCI + IaC/Security\nchecks", COLORS["storage"], fs=11)

    arrow((17, 78), (22, 78))
    arrow((10, 72), (10, 64))
    arrow((17, 58), (22, 72))
    arrow((46, 78), (51, 78))
    arrow((46, 73), (72, 76))
    arrow((46, 69), (72, 60))

    arrow((43, 45), (46, 43), dashed=True)
    arrow((43, 44), (46, 25), dashed=True)
    arrow((62, 43), (85, 43), dashed=True)
    arrow((62, 25), (85, 25), dashed=True)
    arrow((62, 43), (67, 35), dashed=True)
    arrow((62, 25), (67, 33), dashed=True)
    arrow((67, 32), (46, 74), dashed=True)

    arrow((17, 22), (22, 22))
    arrow((40, 22), (46, 40), dashed=True)
    arrow((40, 22), (46, 22), dashed=True)

    ax.text(4, 95, "MyWay AWS Architecture (Current implementation as of September 17, 2026)", fontsize=16, fontweight="bold")
    ax.text(4, 92, "Solid arrows: synchronous request flow. Dashed arrows: scheduled data refresh flow.", fontsize=11)

    legend_y = 6
    for i, (k, label) in enumerate([
        ("app", "User-facing app"),
        ("ai", "AI/data services"),
        ("storage", "Storage/security"),
        ("optional", "Optional scheduled feeds"),
    ]):
        x = 4 + i * 24
        ax.add_patch(FancyBboxPatch((x, legend_y), 20, 4, boxstyle="round,pad=0.02", linewidth=1, edgecolor="#333", facecolor=COLORS[k]))
        ax.text(x + 10, legend_y + 2, label, ha="center", va="center", fontsize=10)

    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def h(doc: Document, level: int, text: str):
    doc.add_heading(text, level=level)


def p(doc: Document, text: str):
    doc.add_paragraph(text)


def bullet(doc: Document, text: str):
    doc.add_paragraph(text, style="List Bullet")


def number(doc: Document, text: str):
    doc.add_paragraph(text, style="List Number")


def make_doc(diagram_path: Path):
    metrics = json.loads(QUESTIONNAIRE_JSON.read_text())

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    hp = section.header.paragraphs[0]
    hp.text = "MyWay Product Requirements Document"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].font.color.rgb = RGBColor(90, 90, 90)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    add_page_number(fp.add_run())

    cover = doc.add_heading("MyWay Product Requirements Document", level=0)
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1 = doc.add_paragraph("Repository: oreoolis/meong-MyWay")
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2 = doc.add_paragraph("Date: September 17, 2026")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3 = doc.add_paragraph("Document status/version: Draft v1.0")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4 = doc.add_paragraph("Purpose: Define MyWay product requirements based on the current implemented system, with clear distinction between implemented, optional, measured, assumed, and future items.")
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h(doc, 1, "Document Revision History")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Version"
    t.rows[0].cells[1].text = "Date"
    t.rows[0].cells[2].text = "Author"
    t.rows[0].cells[3].text = "Summary"
    row = t.add_row().cells
    row[0].text = "v1.0"
    row[1].text = "2026-09-17"
    row[2].text = "GitHub Copilot Coding Agent"
    row[3].text = "Initial PRD generated from repository implementation evidence"

    doc.add_page_break()

    h(doc, 1, "Investor Pitch Q&A Preparation")
    qa = [
        ("How does the questionnaire affect accuracy? Is there any difference with or without it?",
         "The repository’s recorded questionnaire-quality evidence is from a synthetic fixture, not a population benchmark. In that fixture, baseline and skip were identical; adding independent AWS-use evidence moved cloud operations from rank 4 to rank 3, governance evidence moved governance from rank 6 to rank 5, and unrelated pastry similarity remained low. This supports directional usefulness, but broader labeled evaluation is still required."),
        ("What orchestrator are you using?",
         "MyWay uses an in-code fixed-DAG orchestrator in src/lib/agents/orchestrator.ts. The sequencing is deterministic: parser and storage first, then planner, then concurrent specialists (improver and advisor), with the swapper deferred to /api/analysis/swap for latency."),
        ("What harness are you using or is being implemented?",
         "Current harnesses include Vitest unit tests (*.test.ts), integration tests (*.itest.ts) with live AWS/Bedrock calls, fixed-fixture Titan ranking evaluation writing docs/questionnaire-quality-results.json, Lambda pure-function tests, and CI checks (secret scanning plus Terraform/Trivy/Checkov/TFLint/validate). A dedicated production-grade continuous agent-evaluation harness is still future work."),
        ("What is the real difference between CareersFinder on MyCareersFuture and this? If the government implements a new iteration of CareersFinder, how would this fare?",
         "MyWay is not a replacement job-search portal; it is a personalized career-planning layer grounded in an uploaded resume plus optional questionnaire evidence, with resume-improvement, skill-gap, course, adjacent-career, coaching, and live-opening attachment paths. Adaptation is supported by data-source abstraction (API wrappers/stores), deterministic matching layers, evidence-grounded pipelines, and modular agent/tool boundaries. This improves adaptability, but does not guarantee competitive protection against future government iterations.")
    ]
    for q, a in qa:
        number(doc, q)
        ans = doc.add_paragraph(f"Answer: {a}")
        ans.paragraph_format.left_indent = Inches(0.25)

    h(doc, 1, "AWS Architecture Diagram")
    doc.add_picture(str(diagram_path), width=Inches(7.25))
    if doc.inline_shapes:
        doc.inline_shapes[-1]._inline.docPr.set(
            "descr",
            "MyWay AWS architecture showing browser, Next.js routes/API handlers, Cognito, Bedrock, S3, DynamoDB, EventBridge, Lambda scrapers, optional external feeds, and Terraform/GitHub Actions boundaries.",
        )
    add_caption(doc, "Figure 1. MyWay AWS architecture (as implemented on September 17, 2026). Optional components are explicitly labeled.")
    p(doc, "Request path (synchronous): Browser -> Next.js routes/API -> Cognito/session checks + S3/DynamoDB + Bedrock. Scheduled path: EventBridge triggers Lambda scrapers, which refresh jobs/course snapshots in S3, then analysis attachments consume those snapshots.")

    h(doc, 1, "Table of Contents")
    add_toc(doc.add_paragraph())

    h(doc, 1, "Executive Summary")
    p(doc, "MyWay is a Singapore-focused career guidance application that combines a single Next.js deployment, Cognito authentication, resume parsing, optional questionnaire refinement, and six cooperating agents to produce personalized career insights. The current implementation is designed for truthful progress reporting, resumable processing, and evidence-grounded outputs. Optional market-feed infrastructure (jobs and courses scrapers) enriches recommendations when deployed.")

    h(doc, 1, "Problem Statement and Opportunity")
    bullet(doc, "Problem: Students and early-career professionals often receive generic advice that does not map to their actual resume evidence.")
    bullet(doc, "Opportunity: Convert one uploaded resume into practical next steps: improve resume quality, identify role fit and gaps, and show current openings and courses in Singapore.")
    bullet(doc, "Market context: MyCareersFuture provides vacancy discovery; MyWay focuses on personalized readiness and transition planning.")

    h(doc, 1, "Product Vision, Goals, and Non-Goals")
    bullet(doc, "Vision: Make career planning evidence-based, transparent, and actionable for Singapore users.")
    bullet(doc, "Goals (implemented): deterministic orchestration, grounded output, resumability, low-friction questionnaire, optional live data enrichment.")
    bullet(doc, "Non-goals (current): guaranteed hiring outcomes, universal labor-market coverage, replacing government job portals.")

    h(doc, 1, "Target Users/Personas and User Needs")
    bullet(doc, "University students and recent graduates: need realistic role pathways, skill-gap clarity, and resume rewrites.")
    bullet(doc, "Career switchers: need adjacent/pivot destination options and practical milestones.")
    bullet(doc, "Early-career professionals in Singapore: need grounded links between current evidence and local roles, openings, and training options.")

    h(doc, 1, "Scope and Feature Requirements")
    bullet(doc, "Landing/authentication: Amplify Cognito flow with email confirmation, SRP sign-in, access + ID token verification, and optional/production TOTP MFA support (docs/auth/cognito.md; src/lib/auth/*).")
    bullet(doc, "Resume upload: PDF/DOCX only, 5 MB max, client/server validation plus byte sniffing, S3 bytes storage, DynamoDB metadata, one active resume per user.")
    bullet(doc, "Resume intake/questionnaire: parser first, questionnaire agent returns 2-3 reviewed questions from profile gaps; user may answer or skip; canonical IDs; strict server-side validation; first submission immutable; replay/idempotency; lease/expiry guards.")
    bullet(doc, "Final evidence embedding: refined summary + validated factual statements + affirmed/disclaimed skills embedded with Titan Embeddings V2.")
    bullet(doc, "Career analysis branches: planner, improver, advisor in main run; deferred swapper branch loaded separately. Advisor = current-industry view; Transitioner = adjacent/pivot view.")
    bullet(doc, "Live openings (optional infra): MyCareersFuture feed via 12-hour EventBridge/Lambda refresh to S3 snapshot, lexical prefilter then embedding rank, calibrated scores, and skill-gap aggregation; no recommendation if title-relatedness threshold fails.")
    bullet(doc, "Course recommendations (optional infra): weekly SSG-WSG/SkillsFuture scraper with precomputed vectors, then gap-to-course matching.")
    bullet(doc, "UX/progress: one route workspace state machine, truthful phase indicators, real tool-thought traces, responsive questionnaire UI with skip/clear behavior and result branch tabs.")
    bullet(doc, "Failure handling: retries for transient failures, replay with submission keys, resumability via stored artifacts and leases, graceful degradation when optional feeds/services are unavailable.")

    h(doc, 1, "End-to-End Workflows")
    h(doc, 2, "1) Sign-up/sign-in/session verification")
    number(doc, "User signs up/signs in via Cognito SRP in the browser.")
    number(doc, "Browser sends access token and ID token to /api/auth/session.")
    number(doc, "Server verifies signature, issuer, token use, app client, sub matching, and verified email.")
    h(doc, 2, "2) Upload and storage")
    number(doc, "User uploads PDF/DOCX.")
    number(doc, "Server validates extension/type/size and byte signature.")
    number(doc, "Resume bytes are stored in S3; metadata is stored in DynamoDB.")
    h(doc, 2, "3) Intake/questionnaire/answer-or-skip")
    number(doc, "POST /api/analysis/intake parses the current resume profile (without final embedding).")
    number(doc, "Questionnaire agent generates and reviews 2-3 factual gap questions.")
    number(doc, "User submits ID-only selections or skips; server validates canonical IDs and intake version.")
    h(doc, 2, "4) Full analysis")
    p(doc, "Flow diagram: Upload -> Intake parse -> Questionnaire (answer/skip) -> Embed -> Planner -> {Improver || Advisor} -> Results fork -> Deferred Swapper")
    number(doc, "Embedding step runs with validated evidence.")
    number(doc, "Fixed-DAG orchestrator stores profile/embedding, runs planner, then specialist fan-out.")
    number(doc, "Analysis artifacts are stored with TTL and resume/lease transaction checks.")
    h(doc, 2, "5) Advisor results")
    number(doc, "Industry advisor searches/scopes current-sector roles, returns matched roles with rationale/gaps.")
    number(doc, "Openings and courses are attached as independent post-processing layers.")
    h(doc, 2, "6) Transitioner results + deferred swapper")
    number(doc, "Client triggers /api/analysis/swap after results fork renders.")
    number(doc, "Swapper reads stored routing/profile/embedding and returns adjacent/pivot destinations.")
    h(doc, 2, "7) Job/course data refresh")
    p(doc, "Scheduled diagram: EventBridge(12h) -> jobs-scraper Lambda -> MyCareersFuture API -> S3 jobs snapshot; EventBridge(7d) -> courses-scraper Lambda -> SSG course API + Titan embeddings -> S3 courses snapshot")
    number(doc, "EventBridge schedules jobs scraper every 12 hours and courses scraper every 7 days.")
    number(doc, "Lambdas archive run outputs then update latest snapshot pointers.")
    h(doc, 2, "8) Retry/replay/replacement/deletion behavior")
    number(doc, "Idempotent replay occurs for identical questionnaire submissions.")
    number(doc, "Resume replacement/deletion invalidates dependent analyses.")
    number(doc, "Lease-based concurrency control prevents conflicting intake completion workers.")

    h(doc, 1, "AI/Agent Design")
    p(doc, "Implemented agents and responsibilities:")
    at = doc.add_table(rows=1, cols=6)
    at.style = "Table Grid"
    headers = ["Agent", "Primary responsibility", "Inputs", "Outputs", "Tools/use mode", "Run location"]
    for i, c in enumerate(headers):
        at.rows[0].cells[i].text = c
    rows = [
        ("1. Resume Parser", "Extract profile; refine summary and embedding text", "resume bytes + optional questionnaire evidence", "profile + embedding vector/text", "Bounded tool loop + emit_result", "Next.js API runtime"),
        ("2. Questionnaire Agent (Context)", "Generate/review 2-3 factual gap questions", "parsed profile", "questionnaire items + usage", "Single-shot JSON calls", "Next.js API runtime"),
        ("3. Career Planner", "Plan trajectory, paths, sector routing keywords", "profile + framework sectors", "career plan + route keywords", "Bounded tool loop (maxTurns=4)", "Next.js API runtime"),
        ("4. Resume Improver", "Targeted critique and rewrites", "profile + plan + original doc", "improvement package", "Bounded tool loop with keyword-demand check", "Next.js API runtime"),
        ("5. Industry Advisor", "Current-industry role fit and advice", "profile + vector + sector keywords", "matched roles/advice", "Bounded tool loop with framework fallback", "Next.js API runtime"),
        ("6. Career Swapper", "Adjacent/pivot destinations in deferred branch", "stored profile/vector/routing", "swap destinations + coach brief", "Bounded tool loop + reasoned fallback", "Next.js API runtime (/api/analysis/swap)"),
    ]
    for r in rows:
        cells = at.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v

    p(doc, "DAG vs bounded loops: orchestration order is fixed in code (deterministic DAG), while several agents run bounded internal tool loops to validate lookups before emitting structured JSON.")
    p(doc, "Bedrock implementation details: Converse API, forced structured JSON tool output (emit_result), Claude Haiku 4.5 reasoning via US inference profile, Titan Embeddings V2 for vectors, temperature 0, explicit token ceilings per agent, and per-run cost tracking.")
    p(doc, "Grounded vs model-generated outputs: sectors/roles/salary bands come from SSG-WSG APIs when available; model-generated narratives are normalized and constrained; unsupported IDs are dropped; optional reasoned fallback is labeled as reasoned basis.")

    h(doc, 1, "AWS Architecture and Deployment")
    bullet(doc, "Single deployment model: one Next.js application hosts both UI and API routes.")
    bullet(doc, "Infrastructure as code: Terraform modules under iac/ provision S3, DynamoDB, optional EventBridge/Lambda scrapers, and IAM policies.")
    bullet(doc, "Least privilege: scoped IAM permissions for Bedrock model invoke, S3 object access, and DynamoDB operations.")
    bullet(doc, "Encryption/retention: uploads bucket uses SSE-KMS (configurable), analyses/resumes tables have PITR, analyses artifacts carry TTL.")
    bullet(doc, "Configuration boundaries: auth variables, storage variables, Bedrock model IDs/regions, and optional jobs/courses feed toggles are environment-driven.")

    h(doc, 1, "Data Model and Contracts (Conceptual)")
    bullet(doc, "Resume bytes + metadata: file object in S3 and one active metadata row/user in DynamoDB resumes table.")
    bullet(doc, "Intake/questionnaire: intake artifact with parsed profile, questionnaire IDs/options, generation usage/version, submission key, lease, expiry.")
    bullet(doc, "Profile/embedding/plan/routing/specialists: separate artifacts in DynamoDB analyses table.")
    bullet(doc, "Job/course snapshots: non-personal snapshot documents in S3 (jobs/latest.json and courses/latest.json).")
    bullet(doc, "TTL/leases: analyses expire via TTL; intake lock and completion lease support idempotent replay and worker fencing.")

    h(doc, 1, "Non-Functional Requirements")
    bullet(doc, "Security/privacy: token validation, per-user authorization, no token echoing, scoped IAM, encrypted storage, secret scanning in CI.")
    bullet(doc, "Performance/latency: deferred swapper, bounded tool turns, snapshot-based enrichments, role/candidate caps.")
    bullet(doc, "Reliability/recovery: transaction guards, retries for transient AWS failures, replay semantics, PITR, and alarmed scrapers.")
    bullet(doc, "Cost: model ceilings, deterministic temperature, embedding budget caps, optional feed deployment controls.")
    bullet(doc, "Accessibility/usability: semantic labeled questionnaire controls, keyboard flow, truthful progress phases.")
    bullet(doc, "Observability/testability/maintainability: Vitest unit/integration split, Lambda pure-function tests, comments documenting trade-offs and failure modes.")

    h(doc, 1, "Accuracy, Evaluation, and Limitations")
    p(doc, "Measured questionnaire quality evidence (synthetic fixture):")
    qt = doc.add_table(rows=1, cols=6)
    qt.style = "Table Grid"
    for i, c in enumerate(["Variant", "Cloud rank", "Cloud cosine", "Reporting cosine", "Governance cosine", "Unrelated (pastry) cosine"]):
        qt.rows[0].cells[i].text = c
    rankings = metrics["rankings"]
    sims = metrics["cosineSimilarities"]
    variants = ["baseline", "skipped", "independent", "governance"]
    labels = {
        "baseline": "Résumé only",
        "skipped": "Skip",
        "independent": "Independent AWS use",
        "governance": "AWS guidance/governance",
    }
    for v in variants:
        row = qt.add_row().cells
        row[0].text = labels[v]
        row[1].text = str(rankings[v].index("cloud") + 1)
        row[2].text = f"{sims[v]['cloud']:.4f}"
        row[3].text = f"{sims[v]['reporting']:.4f}"
        row[4].text = f"{sims[v]['governance']:.4f}"
        row[5].text = f"{sims[v]['unrelated']:.4f}"
    add_caption(doc, "Table 1. Questionnaire quality fixture results from docs/questionnaire-quality-results.json")

    p(doc, "Interpretation required for investor context: baseline and skip were identical for the synthetic fixture; independent AWS-use evidence moved cloud operations from rank 4 to rank 3; governance evidence moved governance from rank 6 to rank 5; unrelated pastry similarity remained low. This is not population-wide validation.")
    p(doc, "Job matching calibration: MyWay uses lexical prefilter + embedding ranking + calibrated score bands. Limits remain: similarity calibration drift, data-source variability, and role-title ambiguity.")
    p(doc, "Known risks/limitations: self-reported questionnaire evidence can be inaccurate, model output can still err, live API outages and schema drift can degrade enrichment, snapshots can become stale, privacy/cost constraints remain, and broad longitudinal validation is still missing.")

    h(doc, 1, "Investor-Ready Differentiation and Competitive Response")
    bullet(doc, "MyWay differentiates by resume-grounded personalization (planning, rewrites, skill-gaps, adjacent transitions, coaching framing) and by combining static profile evidence with optional live openings/courses.")
    bullet(doc, "Compared to CareersFinder-like discovery tools, MyWay emphasizes candidate-specific preparedness and practical next actions, not only listing jobs.")
    bullet(doc, "If government tools iterate rapidly, MyWay can adapt through modular API wrappers, deterministic ranking layers, and replaceable agent/tool boundaries, but no guaranteed moat is claimed.")

    h(doc, 1, "Metrics and Success Criteria")
    mt = doc.add_table(rows=1, cols=3)
    mt.style = "Table Grid"
    mt.rows[0].cells[0].text = "Metric"
    mt.rows[0].cells[1].text = "Current status"
    mt.rows[0].cells[2].text = "Type"
    metric_rows = [
        ("Activation (sign-up to first upload)", "Suggested", "Suggested"),
        ("Upload completion rate", "Suggested", "Suggested"),
        ("Questionnaire complete/skip split", "Suggested", "Suggested"),
        ("Time to first usable results", "Partially measured in code comments/tests", "Partially measured"),
        ("Result usefulness/user feedback", "Suggested", "Suggested"),
        ("Groundedness (framework-vs-reasoned ratio)", "Available via basis fields", "Currently measurable"),
        ("Opening click-through and course engagement", "Suggested", "Suggested"),
        ("Retry/failure rates", "Suggested", "Suggested"),
        ("Cost per analysis", "Currently tracked in run cost output", "Currently measurable"),
        ("Data freshness (snapshot age)", "Currently measurable from snapshot metadata", "Currently measurable"),
        ("Safety/complaint signals", "Suggested", "Suggested"),
    ]
    for r in metric_rows:
        cells = mt.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = r

    h(doc, 1, "Risks, Assumptions, Dependencies, and Open Questions")
    bullet(doc, "Assumption: users provide truthful and sufficiently detailed resume/questionnaire evidence.")
    bullet(doc, "Dependency: optional SSG-WSG and MyCareersFuture feeds remain reachable and schema-compatible.")
    bullet(doc, "Risk: model token ceilings and cost constraints can trade off output depth.")
    bullet(doc, "Open question: broad, labeled, longitudinal evaluation harness for production-grade quality tracking.")

    h(doc, 1, "Release/Deployment/Runbook Overview")
    number(doc, "Provision infrastructure via Terraform under iac/ with optional scrapers enabled as needed.")
    number(doc, "Configure environment variables for Cognito, storage, Bedrock, and optional feeds.")
    number(doc, "Deploy Next.js app and verify auth/session, upload, intake, analysis, and swap endpoints.")
    number(doc, "Monitor Lambda error alarms and snapshot freshness for jobs/courses feeds.")

    h(doc, 1, "Glossary (for non-specialists)")
    gt = doc.add_table(rows=1, cols=2)
    gt.style = "Table Grid"
    gt.rows[0].cells[0].text = "Term"
    gt.rows[0].cells[1].text = "Plain-English meaning"
    glossary = [
        ("SRP", "A secure sign-in method where the password is not sent in plain form."),
        ("JWT token", "A signed login token proving who the user is and what session is active."),
        ("MFA/TOTP", "A second login factor using one-time numeric codes from an authenticator app."),
        ("Embedding", "A numeric vector representation of text used for similarity matching."),
        ("DAG", "A fixed step order where outputs flow forward without circular loops."),
        ("Cosine similarity", "A score of how semantically close two vectors/texts are."),
        ("TTL", "A data expiry timestamp after which old records are automatically removed."),
        ("PITR", "Point-in-time recovery for restoring DynamoDB tables after accidental data loss."),
    ]
    for term, desc in glossary:
        row = gt.add_row().cells
        row[0].text = term
        row[1].text = desc

    h(doc, 1, "Appendix: Repository-to-Requirement Traceability")
    tr = doc.add_table(rows=1, cols=4)
    tr.style = "Table Grid"
    tr.rows[0].cells[0].text = "Requirement area"
    tr.rows[0].cells[1].text = "Repository evidence"
    tr.rows[0].cells[2].text = "Classification"
    tr.rows[0].cells[3].text = "Notes"
    trace_rows = [
        ("Auth/session", "docs/auth/cognito.md; src/lib/auth/client.ts; src/lib/auth/server.ts; src/app/api/auth/session/route.ts", "Implemented", "SRP + token validation + optional/production MFA path"),
        ("Upload validation/storage", "src/lib/resume/file-policy.ts; src/app/api/resume/route.ts; src/lib/resume/store.ts", "Implemented", "PDF/DOCX + 5 MB + byte sniff + S3+Dynamo"),
        ("Intake/questionnaire", "src/lib/resume/intake.ts; questionnaire.ts; intake-store.ts; src/app/api/analysis/intake/route.ts", "Implemented", "2-3 questions, immutable first submission, leases/idempotency"),
        ("Orchestrator", "src/lib/agents/orchestrator.ts", "Implemented", "Fixed DAG + deferred swapper"),
        ("Agent internals", "src/lib/agents/{resume-parser,career-planner,resume-improver,industry-advisor,career-swapper,resume-context}.ts", "Implemented", "Six agents, mixed grounded/reasoned tiers"),
        ("Bedrock transport", "src/lib/bedrock/reason.ts; embeddings.ts; docs/adr/0002-claude-haiku-for-grounded-advice.md", "Implemented", "Converse tool-use JSON and Titan embeddings"),
        ("Live openings", "src/lib/jobs/*; lambda/jobs-scraper/*; iac/jobs.tf", "Optional", "Enabled only when jobs scraper infra is deployed"),
        ("Course recommendations", "src/lib/courses/*; lambda/courses-scraper/*; iac/courses.tf", "Optional", "Enabled only when courses scraper infra + credentials are deployed"),
        ("IaC/security checks", "iac/*.tf; .github/workflows/verify-infra.yml; .github/workflows/secret-scan.yml", "Implemented", "CI IaC and secret scanning gates"),
        ("Broad production evaluation harness", "vitest.config.mts; *.itest.ts; docs/questionnaire-quality-results.json", "Future", "Current tests exist; broader labeled continuous harness still needed"),
    ]
    for row_data in trace_rows:
        cells = tr.add_row().cells
        for i, v in enumerate(row_data):
            cells[i].text = v

    doc.save(DOC_PATH)


def main():
    draw_architecture(DIAGRAM_PATH)
    make_doc(DIAGRAM_PATH)
    print(f"Wrote {DIAGRAM_PATH}")
    print(f"Wrote {DOC_PATH}")


if __name__ == "__main__":
    main()
